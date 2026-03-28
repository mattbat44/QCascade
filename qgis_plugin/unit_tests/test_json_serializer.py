"""
Tests for JSON serialization wrapper.

@author: GitHub Copilot
"""

import os
import sys
import tempfile
from pathlib import Path

import numpy as np
import pytest

# Add source (src) folder in the python path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))

from json_serializer import save_to_json, load_from_json


def test_save_and_load_simple_dict():
    """Test saving and loading a simple dictionary."""
    data = {
        'key1': 'value1',
        'key2': 42,
        'key3': [1, 2, 3]
    }
    
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'test.json'
        save_to_json(data, filepath)
        loaded_data = load_from_json(filepath)
        
        assert loaded_data == data


def test_save_and_load_numpy_arrays():
    """Test saving and loading numpy arrays with various dtypes."""
    data = {
        'array_float32': np.array([[1.0, 2.0], [3.0, 4.0]], dtype=np.float32),
        'array_float64': np.array([1.1, 2.2, 3.3], dtype=np.float64),
        'array_int': np.array([1, 2, 3], dtype=np.int32),
    }
    
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'test.json'
        save_to_json(data, filepath)
        loaded_data = load_from_json(filepath)
        
        # Check all arrays are present and equal
        assert set(loaded_data.keys()) == set(data.keys())
        for key in data.keys():
            np.testing.assert_array_equal(loaded_data[key], data[key])
            assert loaded_data[key].dtype == data[key].dtype


def test_save_and_load_3d_array():
    """Test saving and loading 3D numpy arrays (common in D-CASCADE)."""
    data = {
        'connectivity': np.random.rand(10, 50, 50).astype(np.float32),
    }
    
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'test.json'
        save_to_json(data, filepath)
        loaded_data = load_from_json(filepath)
        
        assert loaded_data['connectivity'].shape == data['connectivity'].shape
        assert loaded_data['connectivity'].dtype == data['connectivity'].dtype
        np.testing.assert_array_almost_equal(loaded_data['connectivity'], data['connectivity'])


def test_save_and_load_mixed_data():
    """Test saving and loading data with mixed types (like data_output)."""
    data = {
        'Simulation parameters': {
            'psi': [-6, -5, -4],
            'ts length': 86400,
            'update slope': False,
            'idx flow': 1,
        },
        'Volume out [m^3]': np.random.rand(100, 50).astype(np.float32),
        'Volume in [m^3]': np.random.rand(100, 50).astype(np.float32),
        'D50 active layer [m]': np.random.rand(100, 50).astype(np.float32),
    }
    
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'test.json'
        save_to_json(data, filepath)
        loaded_data = load_from_json(filepath)
        
        # Check nested dict
        assert loaded_data['Simulation parameters'] == data['Simulation parameters']
        
        # Check arrays
        for key in ['Volume out [m^3]', 'Volume in [m^3]', 'D50 active layer [m]']:
            assert loaded_data[key].shape == data[key].shape
            assert loaded_data[key].dtype == data[key].dtype
            np.testing.assert_array_almost_equal(loaded_data[key], data[key], decimal=5)


def test_directory_creation():
    """Test that parent directories are created if they don't exist."""
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'subdir' / 'nested' / 'test.json'
        data = {'key': 'value'}
        
        save_to_json(data, filepath)
        assert filepath.exists()
        
        loaded_data = load_from_json(filepath)
        assert loaded_data == data


def test_numpy_scalar_types():
    """Test that numpy scalar types are properly handled."""
    data = {
        'int_val': np.int32(42),
        'float_val': np.float32(3.14),
        'bool_val': np.bool_(True),
    }
    
    with tempfile.TemporaryDirectory() as tmpdir:
        filepath = Path(tmpdir) / 'test.json'
        save_to_json(data, filepath)
        loaded_data = load_from_json(filepath)
        
        # After round-trip, numpy scalars become Python types
        assert loaded_data['int_val'] == 42
        assert abs(loaded_data['float_val'] - 3.14) < 0.01
        assert loaded_data['bool_val'] is True


# ---------------------------------------------------------------------------
# Tests for results_to_spreadsheet
# ---------------------------------------------------------------------------

import pytest

pytest.importorskip("pandas")
pytest.importorskip("openpyxl")

from json_serializer import results_to_spreadsheet  # noqa: E402  (after importorskip)


def test_spreadsheet_creates_file():
    """results_to_spreadsheet creates an .xlsx file."""
    data_output = {
        'Volume out [m^3]': np.random.rand(10, 5).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path)
        assert xlsx_path.exists()


def test_spreadsheet_sheet_per_2d_variable():
    """Each 2-D variable produces exactly one sheet."""
    import openpyxl

    data_output = {
        'Volume out [m^3]': np.random.rand(10, 5).astype(np.float32),
        'Volume in [m^3]': np.random.rand(10, 5).astype(np.float32),
        'Sediment budget [m^3]': np.random.rand(10, 5).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path)

        wb = openpyxl.load_workbook(xlsx_path)
        # Three variables → three sheets
        assert len(wb.sheetnames) == 3


def test_spreadsheet_sheet_per_3d_class():
    """A 3-D variable (time × reach × class) produces one sheet per class."""
    import openpyxl

    n_classes = 4
    data_output = {
        'Volume out per grain sizes [m^3]': np.random.rand(10, 5, n_classes).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path)

        wb = openpyxl.load_workbook(xlsx_path)
        assert len(wb.sheetnames) == n_classes


def test_spreadsheet_reach_ids_as_columns():
    """Custom reach IDs appear as column headers in the sheet."""
    import openpyxl

    reach_ids = [10, 20, 30]
    data_output = {
        'Volume out [m^3]': np.ones((5, 3), dtype=np.float64),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path, reach_ids=reach_ids)

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        # Row 1 = header; columns B, C, D should carry "Reach 10", "Reach 20", "Reach 30"
        headers = [ws.cell(row=1, column=c).value for c in range(2, 5)]
        assert headers == ["Reach 10", "Reach 20", "Reach 30"]


def test_spreadsheet_values_correct():
    """Values written to the sheet match the original numpy array."""
    import openpyxl

    arr = np.arange(12, dtype=np.float64).reshape(3, 4)
    data_output = {'Test variable': arr}
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path)

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        # Data starts at row 2 (row 1 = header), column 2 (column 1 = index)
        for r in range(arr.shape[0]):
            for c in range(arr.shape[1]):
                cell_val = ws.cell(row=r + 2, column=c + 2).value
                assert cell_val == pytest.approx(arr[r, c])


def test_spreadsheet_with_extended_output():
    """Extended output variables are also written to the workbook."""
    import openpyxl

    data_output = {
        'Volume out [m^3]': np.random.rand(10, 3).astype(np.float32),
    }
    extended_output = {
        'Qbi_dep [m^3]': np.random.rand(10, 3).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path, extended_output=extended_output)

        wb = openpyxl.load_workbook(xlsx_path)
        # One sheet from data_output + one from extended_output
        assert len(wb.sheetnames) == 2


def test_spreadsheet_non_array_keys_skipped():
    """Non-array entries (e.g., 'Simulation parameters') are silently skipped."""
    import openpyxl

    data_output = {
        'Simulation parameters': {'psi': [-6, -5], 'ts_length': 86400},
        'Volume out [m^3]': np.random.rand(5, 3).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / 'results.xlsx'
        results_to_spreadsheet(data_output, xlsx_path)

        wb = openpyxl.load_workbook(xlsx_path)
        # Only the array variable produces a sheet
        assert len(wb.sheetnames) == 1


def test_spreadsheet_suffix_added_automatically():
    """If the given path has no .xlsx extension, it is added automatically."""
    data_output = {
        'Volume out [m^3]': np.random.rand(5, 3).astype(np.float32),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        no_ext_path = Path(tmpdir) / 'results'
        results_to_spreadsheet(data_output, no_ext_path)
        assert (Path(tmpdir) / 'results.xlsx').exists()


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
